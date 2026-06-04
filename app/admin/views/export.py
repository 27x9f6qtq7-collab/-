from io import BytesIO
from flask import render_template, request, send_file, current_app
from flask_login import login_required
from dateutil.parser import parse
import openpyxl
from openpyxl.chart import LineChart, BarChart, PieChart, Reference
from collections import defaultdict
from datetime import timedelta
from app import db
from app.models import Order, User, Item, AdminLog
from app.utils import admin_required, utc_to_local
from app.admin import bp

# 导出中心主页
@bp.route('/export-center')
@login_required
@admin_required
def export_center():
    return render_template('admin/export_center.html')

# 导出订单
@bp.route('/export-orders')
@login_required
@admin_required
def export_orders():
    status_filter = request.args.get('status', '')
    query = Order.query
    if status_filter:
        query = query.filter_by(status=status_filter)
    orders = query.order_by(Order.created_at.desc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "订单列表"
    ws.append(['订单ID', '订单号', '商品名称', '买家学号', '金额', '状态', '下单时间（北京时间）'])
    for order in orders:
        buyer = User.query.get(order.buyer_id)
        local_time = utc_to_local(order.created_at)
        ws.append([
            order.id,
            order.order_number,
            order.item.title,
            buyer.student_id,
            order.total_price,
            order.status,
            local_time.strftime('%Y-%m-%d %H:%M:%S') if local_time else ''
        ])
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[col_letter].width = adjusted_width
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name='orders.xlsx', as_attachment=True)

# 导出用户
@bp.route('/export-users')
@login_required
@admin_required
def export_users():
    users = User.query.all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '用户列表'
    ws.append(['ID', '学号', '邮箱', '管理员', '注册时间（北京时间）'])
    for u in users:
        local_time = utc_to_local(u.created_at)
        ws.append([
            u.id,
            u.student_id,
            u.email or '',
            '是' if u.is_admin else '否',
            local_time.strftime('%Y-%m-%d %H:%M:%S') if local_time else ''
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name='users.xlsx', as_attachment=True)

# 导出商品
@bp.route('/export-items')
@login_required
@admin_required
def export_items():
    items = Item.query.all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '商品列表'
    ws.append(['ID', '标题', '描述', '价格', '状态', '卖家学号', '分类', '创建时间（北京时间）'])
    for i in items:
        local_time = utc_to_local(i.created_at)
        ws.append([
            i.id,
            i.title,
            i.description,
            i.price,
            i.status,
            i.seller.student_id,
            i.category.name if i.category else '',
            local_time.strftime('%Y-%m-%d %H:%M:%S') if local_time else ''
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name='items.xlsx', as_attachment=True)

# 按日期范围导出订单
@bp.route('/export-orders-by-date', methods=['GET', 'POST'])
@login_required
@admin_required
def export_orders_by_date():
    if request.method == 'POST':
        start_date = parse(request.form.get('start_date')).date()
        end_date = parse(request.form.get('end_date')).date()
        orders = Order.query.filter(db.func.date(Order.created_at).between(start_date, end_date)).all()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f'订单_{start_date}_{end_date}'
        ws.append(['订单ID', '订单号', '商品名称', '买家学号', '金额', '状态', '下单时间（北京时间）'])
        for order in orders:
            buyer = User.query.get(order.buyer_id)
            local_time = utc_to_local(order.created_at)
            ws.append([
                order.id,
                order.order_number,
                order.item.title,
                buyer.student_id,
                order.total_price,
                order.status,
                local_time.strftime('%Y-%m-%d %H:%M:%S') if local_time else ''
            ])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, download_name=f'orders_{start_date}_{end_date}.xlsx', as_attachment=True)
    return render_template('admin/export_form.html')

# 导出统计报表（含图表）
@bp.route('/export-stats', methods=['GET', 'POST'])
@login_required
@admin_required
def export_stats():
    if request.method == 'POST':
        start_date = parse(request.form.get('start_date')).date()
        end_date = parse(request.form.get('end_date')).date()
        orders = Order.query.filter(db.func.date(Order.created_at).between(start_date, end_date)).all()
        daily_amount = defaultdict(float)
        daily_count = defaultdict(int)
        item_sales = defaultdict(int)
        category_sales = defaultdict(int)
        # 使用本地日期（北京时间）聚合
        for order in orders:
            local_dt = utc_to_local(order.created_at)
            date_str = local_dt.strftime('%Y-%m-%d')
            daily_amount[date_str] += order.total_price
            daily_count[date_str] += 1
            item_sales[order.item_id] += 1
            cat_name = order.item.category.name if order.item.category else '未分类'
            category_sales[cat_name] += 1
        sorted_dates = sorted(daily_amount.keys())
        top_items = sorted(item_sales.items(), key=lambda x: x[1], reverse=True)[:10]
        top_item_objs = [Item.query.get(iid) for iid, _ in top_items]
        wb = openpyxl.Workbook()
        # 订单趋势
        ws_trend = wb.active
        ws_trend.title = "订单趋势"
        ws_trend.append(["日期（北京时间）", "订单金额(¥)", "订单数量"])
        for date in sorted_dates:
            ws_trend.append([date, daily_amount[date], daily_count[date]])
        chart1 = LineChart()
        chart1.title = "每日订单金额趋势"
        chart1.y_axis.title = "金额(¥)"
        chart1.x_axis.title = "日期"
        data = Reference(ws_trend, min_col=2, min_row=1, max_col=2, max_row=len(sorted_dates)+1)
        cats = Reference(ws_trend, min_col=1, min_row=2, max_row=len(sorted_dates)+1)
        chart1.add_data(data, titles_from_data=True)
        chart1.set_categories(cats)
        ws_trend.add_chart(chart1, "E2")
        # 热销商品
        ws_items = wb.create_sheet("热销商品")
        ws_items.append(["商品ID", "商品标题", "销量"])
        for (iid, cnt), item in zip(top_items, top_item_objs):
            ws_items.append([iid, item.title, cnt])
        chart2 = BarChart()
        chart2.title = "热销商品TOP10"
        data2 = Reference(ws_items, min_col=3, min_row=1, max_col=3, max_row=len(top_items)+1)
        cats2 = Reference(ws_items, min_col=2, min_row=2, max_row=len(top_items)+1)
        chart2.add_data(data2, titles_from_data=True)
        chart2.set_categories(cats2)
        ws_items.add_chart(chart2, "E2")
        # 分类占比
        ws_cat = wb.create_sheet("商品分类占比")
        ws_cat.append(["分类", "销量"])
        for cat, cnt in category_sales.items():
            ws_cat.append([cat, cnt])
        chart3 = PieChart()
        data3 = Reference(ws_cat, min_col=2, min_row=1, max_col=2, max_row=len(category_sales)+1)
        labels3 = Reference(ws_cat, min_col=1, min_row=2, max_row=len(category_sales)+1)
        chart3.add_data(data3, titles_from_data=True)
        chart3.set_categories(labels3)
        chart3.title = "商品分类销量占比"
        ws_cat.add_chart(chart3, "D2")
        # 订单明细
        ws_detail = wb.create_sheet("订单明细")
        ws_detail.append(['订单ID', '订单号', '商品', '买家学号', '金额', '状态', '下单时间（北京时间）'])
        for order in orders:
            buyer = User.query.get(order.buyer_id)
            local_time = utc_to_local(order.created_at)
            ws_detail.append([
                order.id,
                order.order_number,
                order.item.title,
                buyer.student_id,
                order.total_price,
                order.status,
                local_time.strftime('%Y-%m-%d %H:%M:%S') if local_time else ''
            ])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, download_name=f'统计报表_{start_date}_{end_date}.xlsx', as_attachment=True)
    return render_template('admin/export_stats_form.html')

# 导出操作日志
@bp.route('/export-logs')
@login_required
@admin_required
def export_logs():
    admin_name = request.args.get('admin_name', '')
    action = request.args.get('action', '')
    start_date = request.args.get('start_date', '')
    end_date = request.args.get('end_date', '')
    query = AdminLog.query
    if admin_name:
        query = query.filter(AdminLog.admin_name.contains(admin_name))
    if action:
        query = query.filter(AdminLog.action == action)
    if start_date:
        query = query.filter(AdminLog.created_at >= parse(start_date))
    if end_date:
        query = query.filter(AdminLog.created_at <= parse(end_date) + timedelta(days=1))
    logs = query.order_by(AdminLog.created_at.desc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "操作日志"
    ws.append(['时间（北京时间）', '管理员', '操作', '目标类型', '目标ID', '详情', 'IP地址'])
    for log in logs:
        local_time = utc_to_local(log.created_at)
        ws.append([
            local_time.strftime('%Y-%m-%d %H:%M:%S'),
            log.admin_name,
            log.action,
            log.target_type or '',
            log.target_id or '',
            log.details or '',
            log.ip_address or ''
        ])
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws.column_dimensions[col_letter].width = adjusted_width
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name='操作日志.xlsx', as_attachment=True)